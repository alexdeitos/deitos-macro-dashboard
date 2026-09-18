from __future__ import annotations

import uuid

from django.core.cache import cache
from django.db import connection
from django.http import JsonResponse
from django.shortcuts import render
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_GET, require_POST

from .models import CapturePoint, CollectionRun
from .services.collector import MarketCollector
from .services.economic_calendar import TradingEconomicsCalendarCollector, calendar_payload
from .services.daytrade import build_daytrade
from .services.news import InvestingNewsCollector, news_payload
from .services.persistence import get_latest_payload, history_payload, persist_payload
from .services.remote_market import remote_market_enabled
import hmac
import os
REFRESH_LOCK_KEY = "market-dashboard:refresh-lock"


@ensure_csrf_cookie
@require_GET
def index(request):
    return render(request, "dashboard/index.html")

@require_GET
def api_public_market_snapshot(request):
    expected_token = os.getenv("PUBLIC_MARKET_API_TOKEN", "")
    received_token = request.headers.get("X-Market-Token", "")

    if not expected_token or not hmac.compare_digest(
        received_token,
        expected_token,
    ):
        return JsonResponse({"detail": "Unauthorized"}, status=401)

    payload = get_latest_payload()
    if payload is None:
        return JsonResponse(
            {"available": False, "detail": "Nenhuma coleta disponível."},
            status=503,
        )

    return JsonResponse(
        payload,
        json_dumps_params={"ensure_ascii": False},
    )
    
@ensure_csrf_cookie
@require_GET
def daytrade(request):
    return render(request, "dashboard/daytrade.html")

@ensure_csrf_cookie
@require_GET
def clean_panel(request):
    return render(request, "dashboard/clean_panel.html")


@require_GET
def api_daytrade(request):
    return JsonResponse(build_daytrade(), json_dumps_params={"ensure_ascii": False})


@require_POST
def api_daytrade_refresh(request):
    return JsonResponse(build_daytrade(force=True), json_dumps_params={"ensure_ascii": False})


@require_GET
def validation(request):
    payload = get_latest_payload()
    return render(request, "dashboard/validation.html", {"payload": payload, "source_status": (payload or {}).get("source_status", {})})


@require_GET
def api_dashboard(request):
    payload = get_latest_payload()
    # No SQLite temporário da Vercel, a primeira instância pode nascer sem dados.
    # Quando o modo remoto está ativo, inicializa automaticamente usando o JSON local.
    if payload is None and remote_market_enabled():
        try:
            _sync_market_collection()
            payload = get_latest_payload()
        except Exception:
            payload = None
    if payload is None:
        return JsonResponse({"available": False, "message": "Ainda não existe uma coleta válida. Use o botão Atualizar.", "data_policy": "Nenhum valor de referência foi usado."}, status=503)
    response = dict(payload)
    response["available"] = True
    response["history"] = history_payload()

    # FRP0 is authoritative from the Excel/Profit CONFIG_CAPTURA capture.
    # Priority: fresh Windows COM bridge (live in-memory Excel) -> current saved
    # workbook -> older database snapshot. This prevents a stale cached workbook
    # from overriding a live RTD value.
    frp0_point = (
        CapturePoint.objects.filter(
            sheet_name="CONFIG_CAPTURA",
            symbol__iexact="FRP0",
            value__isnull=False,
            metadata__source="profit_excel_com",
        )
        .order_by("-observed_at", "-id")
        .first()
    )

    workbook_frp0 = None
    workbook_frp0_cell = None
    workbook_frp0_mtime = None
    try:
        from .services.capture_import import configured_live_workbook_path, _number
        from openpyxl import load_workbook

        workbook_path = configured_live_workbook_path()
        if workbook_path.exists():
            workbook_frp0_mtime = workbook_path.stat().st_mtime
            wb = load_workbook(workbook_path, read_only=True, data_only=True)
            try:
                ws = wb["CONFIG_CAPTURA"]
                for excel_row, row in enumerate(ws.iter_rows(min_row=10, values_only=True), start=10):
                    symbol = str(row[0] or "").strip().upper() if row else ""
                    if symbol != "FRP0":
                        continue
                    value = _number(row[1] if len(row) > 1 else None)
                    if value is not None:
                        workbook_frp0 = value
                        workbook_frp0_cell = f"B{excel_row}"
                        break
            finally:
                wb.close()
    except Exception:
        workbook_frp0 = None
        workbook_frp0_cell = None
        workbook_frp0_mtime = None

    chosen_value = None
    chosen_at = None
    chosen_source = None

    # A recent COM bridge sample represents the actual in-memory RTD value.
    if frp0_point is not None:
        age = (timezone.now() - frp0_point.observed_at).total_seconds()
        if 0 <= age <= 90:
            chosen_value = frp0_point.value
            chosen_at = frp0_point.observed_at
            chosen_source = "Excel/Profit via ponte"
            frp0_point = None  # prevent fallback selection below

    # When there is no fresh live bridge sample, use the workbook itself.
    if chosen_value is None and workbook_frp0 is not None:
        from datetime import datetime as _datetime, timezone as _dt_timezone
        from django.utils import timezone as _timezone
        chosen_value = workbook_frp0
        if workbook_frp0_mtime is not None:
            chosen_at = _datetime.fromtimestamp(
                workbook_frp0_mtime,
                tz=_dt_timezone.utc,
            ).astimezone(_timezone.get_current_timezone())
        chosen_source = "COTACOES.xlsm / CONFIG_CAPTURA"

    # Last-resort database fallback keeps the dashboard useful when Excel has
    # not yet been saved/mounted in the container.
    if chosen_value is None:
        fallback_point = (
            CapturePoint.objects.filter(
                sheet_name="CONFIG_CAPTURA",
                symbol__iexact="FRP0",
                value__isnull=False,
            )
            .order_by("-observed_at", "-id")
            .first()
        )
        if fallback_point is not None:
            chosen_value = fallback_point.value
            chosen_at = fallback_point.observed_at
            chosen_source = "CONFIG_CAPTURA (última captura salva)"

    response["profit_excel"] = {
        "frp0_points": chosen_value,
        "frp0_observed_at": chosen_at.isoformat() if chosen_at else None,
        "frp0_source": chosen_source,
        "frp0_sheet": "CONFIG_CAPTURA",
        "frp0_cell": workbook_frp0_cell or "B37",
    }
    return JsonResponse(response)


def _sync_market_collection() -> dict:
    run = CollectionRun.objects.create(task_id=f"vercel-{uuid.uuid4().hex[:12]}")
    try:
        payload = MarketCollector().collect()
        persist_payload(payload, run)
        run.status = CollectionRun.Status.SUCCESS if payload.get("is_complete") else CollectionRun.Status.PARTIAL
        run.source_status = payload.get("source_status", {})
        run.finished_at = timezone.now()
        run.save(update_fields=["status", "source_status", "finished_at"])
        return {"status": run.status, "collected_at": payload["collected_at"]}
    except Exception as exc:
        run.status = CollectionRun.Status.FAILED
        run.error = str(exc)
        run.finished_at = timezone.now()
        run.save(update_fields=["status", "error", "finished_at"])
        raise


@require_POST
def api_refresh(request):
    if not cache.add(REFRESH_LOCK_KEY, "1", timeout=60):
        return JsonResponse({"accepted": False, "message": "Já existe uma atualização recente em andamento."}, status=429)
    try:
        result = _sync_market_collection()
        return JsonResponse({"accepted": True, "completed": True, "result": result})
    except Exception as exc:
        return JsonResponse({"accepted": False, "message": f"Falha na coleta: {exc}"}, status=503)
    finally:
        cache.delete(REFRESH_LOCK_KEY)


@require_GET
def api_task_status(request, task_id: str):
    return JsonResponse({"task_id": task_id, "state": "SUCCESS", "ready": True, "result": {"status": "completed"}})


@require_GET
def api_news(request):
    try:
        limit = int(request.GET.get("limit", "50"))
    except ValueError:
        limit = 50
    return JsonResponse(news_payload(limit=limit, market=request.GET.get("market", ""), category=request.GET.get("category", "")))


@require_POST
def api_refresh_news(request):
    try:
        result = InvestingNewsCollector().collect()
        return JsonResponse({"accepted": True, "completed": True, "result": result})
    except Exception as exc:
        return JsonResponse({"accepted": False, "message": f"Não foi possível atualizar as notícias: {exc}"}, status=503)


@require_GET
def api_calendar(request):
    try:
        days = int(request.GET.get("days", "7"))
    except ValueError:
        days = 7
    try:
        importance = int(request.GET.get("importance", "1"))
    except ValueError:
        importance = 1
    return JsonResponse(calendar_payload(days=days, country=request.GET.get("country", ""), min_importance=importance))


@require_POST
def api_refresh_calendar(request):
    try:
        result = TradingEconomicsCalendarCollector().collect()
        return JsonResponse({"accepted": True, "completed": True, "result": result})
    except Exception as exc:
        return JsonResponse({"accepted": False, "message": f"Não foi possível atualizar o calendário: {exc}"}, status=503)


@require_GET
def api_raw_snapshot(request):
    payload = get_latest_payload()
    if payload is None:
        return JsonResponse({"available": False}, status=404)
    return JsonResponse(payload)


@require_GET
def health(request):
    checks = {"database": False, "cache": False}
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1")
            checks["database"] = cursor.fetchone()[0] == 1
    except Exception:
        pass
    try:
        cache.set("healthcheck", "ok", timeout=5)
        checks["cache"] = cache.get("healthcheck") == "ok"
    except Exception:
        pass
    healthy = all(checks.values())
    return JsonResponse({"status": "ok" if healthy else "degraded", "checks": checks}, status=200 if healthy else 503)
