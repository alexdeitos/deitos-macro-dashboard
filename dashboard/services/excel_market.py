from __future__ import annotations

from typing import Any, Iterable

from django.utils import timezone

from .capture_import import _number, _read_live_config_rows, configured_live_workbook_path
from dashboard.models import CapturePoint


# The dashboard accepts the common aliases people use when naming the
# instruments in CONFIG_CAPTURA. The canonical keys below are what the API
# exposes to the front-end.
SYMBOL_ALIASES: dict[str, tuple[str, ...]] = {
    "IBOV": (
        "IBOV",
        "IBOVESPA",
        "IBOVFUT",
        "IBOVESPAFUT",
    ),
    "WINFUT": (
        "WINFUT",
        "WIN",
    ),
    "SP500_FUT": (
        "SP500_FUT",
        "SP500FUT",
        "ES1!",
        "ES",
    ),
}


def _alias_rank(symbol: str, aliases: tuple[str, ...]) -> int:
    normalized = str(symbol or "").strip().upper()
    try:
        return aliases.index(normalized)
    except ValueError:
        return len(aliases)


def _pick_alias_row(rows: Iterable[dict[str, Any]], aliases: tuple[str, ...]) -> dict[str, Any] | None:
    candidates = [
        row for row in rows
        if str(row.get("symbol") or "").strip().upper() in aliases
        and row.get("value") is not None
    ]
    if not candidates:
        return None
    candidates.sort(key=lambda row: (_alias_rank(row.get("symbol", ""), aliases), -int(row.get("row_number", 0))))
    return candidates[0]


def _fresh_bridge_rows(aliases_by_key: dict[str, tuple[str, ...]], max_age_seconds: int = 90) -> dict[str, dict[str, Any]]:
    latest = (
        CapturePoint.objects.filter(
            metadata__source="profit_excel_com",
            sheet_name="CONFIG_CAPTURA",
        )
        .order_by("-observed_at", "-id")
        .first()
    )
    if latest is None:
        return {}

    age = (timezone.now() - latest.observed_at).total_seconds()
    if age < 0 or age > max_age_seconds:
        return {}

    rows = list(
        CapturePoint.objects.filter(
            metadata__source="profit_excel_com",
            sheet_name="CONFIG_CAPTURA",
            observed_at=latest.observed_at,
            value__isnull=False,
        ).order_by("id")
    )

    result: dict[str, dict[str, Any]] = {}
    for key, aliases in aliases_by_key.items():
        row = _pick_alias_row(
            [
                {
                    "symbol": item.symbol,
                    "value": item.value,
                    "change_percent": item.change_percent,
                    "trades": item.trades,
                    "volume": item.volume,
                }
                for item in rows
            ],
            aliases,
        )
        if row is not None:
            result[key] = {
                **row,
                "observed_at": latest.observed_at.isoformat(),
                "source": "Excel/Profit via ponte",
                "sheet": "CONFIG_CAPTURA",
            }
    return result


def _workbook_rows() -> list[dict[str, Any]]:
    path = configured_live_workbook_path()
    if not path.exists():
        return []

    rows: list[dict[str, Any]] = []
    for index, row in enumerate(_read_live_config_rows(path), start=10):
        row = dict(row)
        row["row_number"] = index
        rows.append(row)
    return rows


def _saved_workbook_values(aliases_by_key: dict[str, tuple[str, ...]]) -> dict[str, dict[str, Any]]:
    try:
        rows = _workbook_rows()
    except Exception:
        return {}
    if not rows:
        return {}

    result: dict[str, dict[str, Any]] = {}
    for key, aliases in aliases_by_key.items():
        row = _pick_alias_row(rows, aliases)
        if row is None:
            continue
        result[key] = {
            "value": float(row["value"]),
            "change_percent": _number(row.get("field_c")),
            "trades": _number(row.get("field_d")),
            "volume": _number(row.get("field_e")),
            "symbol": row.get("symbol"),
            "row_number": row.get("row_number"),
            "observed_at": None,
            "source": "COTACOES.xlsm / CONFIG_CAPTURA",
            "sheet": "CONFIG_CAPTURA",
        }
    return result


def _database_fallback(aliases_by_key: dict[str, tuple[str, ...]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for key, aliases in aliases_by_key.items():
        points = (
            CapturePoint.objects.filter(
                sheet_name="CONFIG_CAPTURA",
                value__isnull=False,
            )
            .order_by("-observed_at", "-id")
        )
        for point in points:
            symbol = str(point.symbol or "").strip().upper()
            if symbol not in aliases:
                continue
            result[key] = {
                "value": point.value,
                "change_percent": point.change_percent,
                "trades": point.trades,
                "volume": point.volume,
                "symbol": point.symbol,
                "row_number": None,
                "observed_at": point.observed_at.isoformat() if point.observed_at else None,
                "source": "CONFIG_CAPTURA (última captura salva)",
                "sheet": "CONFIG_CAPTURA",
            }
            break
    return result


def read_workbook_cell(cell: str, path=None, expected_symbol: str | None = None) -> dict[str, Any] | None:
    """Read one authoritative cell from the configured COTACOES workbook.

    This is intentionally independent from CONFIG_CAPTURA symbol matching so
    a dashboard card can be tied to an exact Excel cell when requested.
    When ``expected_symbol`` is supplied it is only an optional guard; callers
    that require a literal cell value should omit it.
    """
    from openpyxl import load_workbook

    workbook_path = path or configured_live_workbook_path()
    if not workbook_path.exists():
        return None
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "CONFIG_CAPTURA" not in wb.sheetnames:
            return None
        ws = wb["CONFIG_CAPTURA"]
        if expected_symbol:
            row = ws[cell].row
            symbol = str(ws.cell(row=row, column=1).value or "").strip().upper()
            if symbol != expected_symbol.strip().upper():
                return None
        value = _number(ws[cell].value)
        if value is None:
            return None
        return {
            "value": value,
            "cell": cell.upper(),
            "source": "COTACOES.xlsm / CONFIG_CAPTURA",
            "sheet": "CONFIG_CAPTURA",
        }
    finally:
        wb.close()


def read_excel_market_snapshot(
    symbols: Iterable[str] = ("IBOV", "WINFUT", "SP500_FUT"),
) -> dict[str, dict[str, Any]]:
    """Read live market values from the Profit/Excel capture in priority order.

    Priority is:
      1. fresh Windows COM bridge sample;
      2. saved COTACOES.xlsm / CONFIG_CAPTURA values;
      3. last database capture from CONFIG_CAPTURA.

    Missing values are left missing. No synthetic market value is created.
    """
    requested = tuple(str(symbol).strip().upper() for symbol in symbols)
    aliases_by_key = {key: SYMBOL_ALIASES[key] for key in requested if key in SYMBOL_ALIASES}

    result = _fresh_bridge_rows(aliases_by_key)
    workbook = _saved_workbook_values(aliases_by_key)
    result.update({key: value for key, value in workbook.items() if key not in result})
    fallback = _database_fallback(aliases_by_key)
    result.update({key: value for key, value in fallback.items() if key not in result})
    return result
