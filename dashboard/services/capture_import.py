from __future__ import annotations

from datetime import datetime
from io import BytesIO
import os
from pathlib import Path
from typing import Any

from django.db import transaction
from django.utils.dateparse import parse_datetime
from django.utils import timezone
from openpyxl import load_workbook

from dashboard.models import CapturePoint


def configured_live_workbook_path() -> Path:
    """Resolve the workbook that Profit/Excel is updating.

    No .env change is required. When CAPTURE_XLSM_PATH is not configured,
    the newest COTACOES*.xlsm/.xlsx file under /app/data is selected so that
    renaming the workbook does not silently disconnect the live radar.
    """
    raw = os.getenv("CAPTURE_XLSM_PATH", "").strip()
    if raw:
        path = Path(raw).expanduser()
        if not path.is_absolute():
            path = Path.cwd() / path
        return path

    data_dir = Path("/app/data")
    candidates = list(data_dir.glob("COTACOES*.xlsm")) + list(data_dir.glob("COTACOES*.xlsx"))
    if candidates:
        return max(candidates, key=lambda item: item.stat().st_mtime)

    return data_dir / "COTACOES.xlsm"


def _read_live_config_rows(path: Path):
    if not path.exists():
        raise FileNotFoundError(f"Arquivo de captura não encontrado: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    diagnostics = {"rows": 0, "numeric": 0, "error_values": 0, "empty": 0}
    try:
        if "CONFIG_CAPTURA" not in wb.sheetnames:
            raise ValueError("Aba CONFIG_CAPTURA não encontrada no arquivo Excel.")
        ws = wb["CONFIG_CAPTURA"]
        for row in ws.iter_rows(min_row=10, values_only=True):
            diagnostics["rows"] += 1
            if not row:
                diagnostics["empty"] += 1
                continue
            symbol = str(row[0] or "").strip().upper()
            if not symbol:
                diagnostics["empty"] += 1
                continue
            raw_value = row[1] if len(row) > 1 else None
            if isinstance(raw_value, str) and raw_value.strip().startswith("#"):
                diagnostics["error_values"] += 1
            value = _number(raw_value)
            if value is None or value <= 0:
                continue
            diagnostics["numeric"] += 1
            yield {
                "symbol": symbol,
                "value": value,
                "field_c": row[2] if len(row) > 2 else None,
                "field_d": row[3] if len(row) > 3 else None,
                "field_e": row[4] if len(row) > 4 else None,
            }
    finally:
        wb.close()


def _fresh_profit_bridge_capture(max_age_seconds: int = 90) -> dict[str, Any] | None:
    """Return the latest Windows/Excel COM bridge capture when it is fresh."""
    cutoff = timezone.now() - timezone.timedelta(seconds=max_age_seconds)
    latest = (
        CapturePoint.objects.filter(
            metadata__source="profit_excel_com",
            observed_at__gte=cutoff,
            sheet_name="CONFIG_CAPTURA",
        )
        .order_by("-observed_at", "-id")
        .first()
    )
    if not latest:
        return None
    return {
        "observed_at": latest.observed_at.isoformat(),
        "asset_count": CapturePoint.objects.filter(
            metadata__source="profit_excel_com",
            sheet_name="CONFIG_CAPTURA",
            observed_at=latest.observed_at,
        ).count(),
    }


@transaction.atomic
def sync_live_workbook() -> dict[str, Any]:
    """Synchronize CONFIG_CAPTURA.

    The Linux/Docker container cannot read Excel's in-memory RTD state. When
    the Windows COM bridge is running, its fresh captures are the authoritative
    live source. Only when no bridge sample is fresh do we fall back to the
    workbook's saved cache.
    """
    bridge = _fresh_profit_bridge_capture()
    if bridge:
        return {
            "status": "success_bridge",
            "source": "profit_excel_com",
            "observed_at": bridge["observed_at"],
            "imported": bridge["asset_count"],
            "message": "Captura ao vivo recebida do Excel/Profit via ponte Windows COM.",
        }

    path = configured_live_workbook_path()
    snapshot_at = timezone.now().replace(second=0, microsecond=0)

    # Prevent duplicate snapshots if Celery beat retries or is briefly restarted.
    if CapturePoint.objects.filter(
        observed_at=snapshot_at, sheet_name="CONFIG_CAPTURA", metadata__auto_sync=True
    ).exists():
        return {
            "status": "skipped",
            "reason": "minute_already_synced",
            "observed_at": snapshot_at.isoformat(),
            "path": str(path),
        }

    reader = _read_live_config_rows(path)
    rows = list(reader)
    if not rows:
        # The workbook can contain RTD formulas with #N/A cached values. This
        # is a meaningful diagnostic: Docker can read the saved cache, but it
        # cannot repair Excel/RTD itself.
        return {
            "status": "empty",
            "observed_at": snapshot_at.isoformat(),
            "path": str(path),
            "message": "Nenhum valor numérico disponível na CONFIG_CAPTURA. Verifique o RTDTrading.RTDServer no Excel/Profit ou use a ponte Windows COM.",
        }

    objects = []
    for row in rows:
        current = float(row["value"])
        # CONFIG_CAPTURA column C is the Profit RTD VAR field and is the
        # authoritative daily/market variation. Never derive variation from
        # the previous Django snapshot: that would turn a daily VAR into a
        # one-minute/one-interval return and can invert the analysis.
        change = _number(row.get("field_c"))
        objects.append(
            CapturePoint(
                observed_at=snapshot_at,
                sheet_name="CONFIG_CAPTURA",
                symbol=row["symbol"][:40],
                value=current,
                change_percent=change,
                metadata={
                    "auto_sync": True,
                    "source": "CONFIG_CAPTURA",
                    "path": str(path),
                    "rtD_field_c": str(row["field_c"]) if row["field_c"] is not None else None,
                    "rtD_field_d": str(row["field_d"]) if row["field_d"] is not None else None,
                    "rtD_field_e": str(row["field_e"]) if row["field_e"] is not None else None,
                    "variation_source": "profit_rtd_var",
                },
            )
        )
    CapturePoint.objects.bulk_create(objects, batch_size=1000)
    # New Excel snapshot must be visible immediately to the radar API.
    try:
        from django.core.cache import cache
        cache.delete("macro-dashboard:index-radar:v2")
    except Exception:
        pass
    try:
        mtime = path.stat().st_mtime
    except OSError:
        mtime = None
    return {
        "status": "success",
        "imported": len(objects),
        "observed_at": snapshot_at.isoformat(),
        "path": str(path),
        "file_mtime": mtime,
    }


def _number(value: Any) -> float | None:
    """Parse Excel/RTD numeric values without corrupting decimal points.

    Accepts numeric cells plus pt-BR strings (1.234,56) and plain decimal
    strings (1234.56). Percent signs are ignored.
    """
    if value in (None, ""):
        return None
    try:
        if isinstance(value, str):
            text = value.strip().replace("%", "").replace("\xa0", "").replace(" ", "")
            if not text:
                return None
            if "," in text and "." in text:
                # Last separator is the decimal separator.
                if text.rfind(",") > text.rfind("."):
                    text = text.replace(".", "").replace(",", ".")
                else:
                    text = text.replace(",", "")
            elif "," in text:
                text = text.replace(",", ".")
            return float(text)
        return float(value)
    except (ValueError, TypeError):
        return None


def _datetime(value: Any):
    if isinstance(value, datetime):
        dt = value
        if timezone.is_naive(dt):
            dt = timezone.make_aware(dt)
        return dt
    if value:
        dt = parse_datetime(str(value))
        if dt:
            return timezone.make_aware(dt) if timezone.is_naive(dt) else dt
    return timezone.now()


@transaction.atomic
def import_workbook(uploaded_file) -> dict[str, int]:
    content = uploaded_file.read()
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb["HISTORICO"] if "HISTORICO" in wb.sheetnames else None
    imported = 0
    skipped = 0

    if ws is None:
        # Fallback: import current grid from the first sheet when the workbook
        # has not yet started recording HISTORICO.
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        for row in rows:
            if not row or not row[0] or not isinstance(row[0], str):
                continue
            symbol = str(row[0]).strip().upper()
            if not symbol or symbol in {"ATIVO", "CÓDIGO", "CODIGO"}:
                continue
            value = _number(row[1] if len(row) > 1 else None)
            change = _number(row[2] if len(row) > 2 else None)
            if value is None and change is None:
                continue
            CapturePoint.objects.create(
                observed_at=timezone.now(), sheet_name=ws.title, symbol=symbol,
                value=value, change_percent=change,
                trades=_number(row[3] if len(row) > 3 else None),
                volume=_number(row[4] if len(row) > 4 else None),
                metadata={"import_mode": "current_grid"},
            )
            imported += 1
        return {"imported": imported, "skipped": skipped}

    headers = [str(x or "").strip().lower() for x in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = {h: i for i, h in enumerate(headers)}

    def col(*names):
        for name in names:
            if name in idx:
                return idx[name]
        return None

    time_i = col("data/hora", "data hora", "timestamp", "observed_at")
    sheet_i = col("aba", "sheet", "sheet_name")
    symbol_i = col("ativo", "symbol", "código", "codigo")
    value_i = col("último", "ultimo", "value")
    change_i = col("variação", "variacao", "change_percent")
    trades_i = col("negócios", "negocios", "trades")
    volume_i = col("volume")

    if symbol_i is None:
        raise ValueError("Aba HISTORICO sem coluna Ativo/Symbol.")

    batch = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        symbol = str(row[symbol_i] or "").strip().upper()
        if not symbol:
            skipped += 1
            continue
        value = _number(row[value_i]) if value_i is not None and value_i < len(row) else None
        change = _number(row[change_i]) if change_i is not None and change_i < len(row) else None
        if value is None and change is None:
            skipped += 1
            continue
        batch.append(
            CapturePoint(
                observed_at=_datetime(row[time_i]) if time_i is not None and time_i < len(row) else timezone.now(),
                sheet_name=str(row[sheet_i] if sheet_i is not None and sheet_i < len(row) else "HISTORICO")[:80],
                symbol=symbol[:40],
                value=value,
                change_percent=change,
                trades=_number(row[trades_i]) if trades_i is not None and trades_i < len(row) else None,
                volume=_number(row[volume_i]) if volume_i is not None and volume_i < len(row) else None,
                metadata={"import_mode": "historico"},
            )
        )
        if len(batch) >= 1000:
            CapturePoint.objects.bulk_create(batch, batch_size=1000)
            imported += len(batch)
            batch.clear()
    if batch:
        CapturePoint.objects.bulk_create(batch, batch_size=1000)
        imported += len(batch)
    return {"imported": imported, "skipped": skipped}
